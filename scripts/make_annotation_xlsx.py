"""Etiketleme_A.xlsx ve Etiketleme_B.xlsx: aynı maddeler, farklı sıra; kılavuz ilk sayfada."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

items = pd.read_csv("out/annotation_items.csv")
LABELS = ["Pozitif", "Negatif", "Nötr", "Belirsiz", "Konu dışı / anlamsız"]
F = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")

GUIDE = [
    ("h1", "İnsan Etiketleme Kılavuzu"),
    ("p", "Bu dosyada {n} kısa yorum var. Her yorumu okuyup 'Etiketleme' sayfasındaki sarı 'Etiket' hücresinden açılır listeyle bir etiket seçin. 'Not' sütunu isteğe bağlıdır."),
    ("p", "İlerleme: aşağıdaki hücre doldurduğunuz madde sayısını gösterir."),
    ("progress", ""),
    ("h2", "Temel ilke"),
    ("p", "Etiket yalnızca metnin ifade ettiği duyguya göre verilir. Yıldız puanı, veri setindeki mevcut etiket veya başka bir bilgi size gösterilmez ve dikkate alınmaz."),
    ("h2", "Etiketler"),
    ("t", "Pozitif|Metin, değerlendirilen ürün, hizmet, film veya deneyim hakkında genel olarak olumlu bir görüş, memnuniyet, beğeni veya tavsiye ifade ediyorsa."),
    ("t", "Negatif|Metin genel olarak memnuniyetsizlik, şikâyet, eleştiri, hayal kırıklığı veya tavsiye etmeme anlamı taşıyorsa."),
    ("t", "Nötr|Metin anlaşılır ve konuyla ilgili, ancak belirgin biçimde olumlu ya da olumsuz bir tutum taşımıyorsa veya çok hafif/değerlendirme yönü zayıf bir tutum ifade ediyorsa."),
    ("t", "Belirsiz|Olumlu veya olumsuz yön güvenilir biçimde belirlenemiyorsa; olumlu ve olumsuz ifadeler birbirini yaklaşık olarak dengeliyorsa; ironi, eksik bağlam veya dilsel belirsizlik nedeniyle karar verilemiyorsa."),
    ("t", "Konu dışı / anlamsız|Metin değerlendirilen nesneyle ilgili bir görüş içermiyorsa, anlamsız/bozuksa veya duygu değerlendirmesine imkân vermeyecek ölçüde ilgisizse."),
    ("p", "Nötr ile Belirsiz aynı değildir. Nötr: anlam açık, duygu yönü nötr. Belirsiz: hangi sınıfa ait olduğuna güvenilir biçimde karar verilemiyor."),
    ("h2", "Karar verme kuralları"),
    ("p", "1. Önce metnin anlaşılır ve konuyla ilgili olup olmadığına bakın. Değilse 'Konu dışı / anlamsız'."),
    ("p", "2. Açık beğeni veya memnuniyet varsa 'Pozitif'; açık şikâyet veya memnuniyetsizlik varsa 'Negatif'."),
    ("p", "3. İlgili ve anlaşılır ama belirgin olumlu/olumsuz değerlendirme yoksa 'Nötr'."),
    ("p", "4. Hem olumlu hem olumsuz ifade varsa baskın yöne bakın; bir yön açıkça baskınsa o etiketi kullanın. Örnek: 'Ürün güzel ama kargo çok geç geldi, yine de tavsiye ederim' → Pozitif."),
    ("p", "5. Olumlu ve olumsuz yönler dengeliyse 'Belirsiz'. Karışık yorumları otomatik olarak 'Nötr' kabul etmeyin."),
    ("p", "6. Olumsuzlamaya dikkat: 'fena değil', 'kötü değil', 'pişman olmadım' olumsuz sözcük içerse de genellikle olumludur."),
    ("p", "7. Yalnızca özellik bildirmek duygu değildir: '500 ml' → Nötr; 'çok küçük, 500 ml ancak' → Negatif."),
    ("p", "8. Fiyat, koku, kalite, performans gibi tek bir özelliğe yönelik açık değerlendirme, metinde başka bir tutum yoksa genel duygu kabul edilir."),
    ("p", "9. Değerlendirmenin asıl nesnesi ürün, hizmet veya filmdir. Kargo, paketleme, satıcı gibi ikincil konulardaki övgü/şikâyet, ürüne ilişkin açık bir değerlendirmenin önüne geçmez. Yalnızca bu konulardan söz eden yorumlar ('hızlı kargo, teşekkürler') kendi yönüne göre etiketlenir."),
    ("h2", "Zor örnekler"),
    ("t", "idare eder|Nötr. Ek olumsuz içerik varsa ('ancak idare eder, bir daha almam') Negatif."),
    ("t", "fena değil|Pozitif (hafif olumlu)."),
    ("t", "güzel ürün|Pozitif."),
    ("t", "uygun fiyat|Pozitif. Yalnızca 'fiyat 500 TL' gibi betimleme Nötr."),
    ("t", "kokusu biraz ağır|Negatif (bir özelliğe yönelik şikâyet)."),
    ("t", "fiyatına göre güzel bir ürün|Pozitif (koşullu ama olumlu)."),
    ("t", "Kalitesi çok iyi ama kargo geç geldi.|Pozitif (kural 9: ürün değerlendirmesi baskın)."),
    ("t", "Görüntüsü güzel ama çok yavaş, sürekli donuyor ve aldığıma pişman oldum.|Negatif."),
    ("t", "Oyunculuk iyi, senaryo zayıf; ne beğendim ne de tamamen kötü buldum.|Belirsiz."),
    ("h2", "Film yorumları"),
    ("p", "Duygu, filmde anlatılan olaylara göre değil, yorumcunun filme ilişkin değerlendirmesine göre belirlenir. 'Film çok acıklı, herkes ölüyor ama kesinlikle izlenmesi gereken harika bir yapım' → Pozitif. Yalnızca konuyu özetleyen, değerlendirme içermeyen metin → Nötr. Bazı film yorumları Türkçe karakter kullanılmadan yazılmıştır (ör. 'begendim'); bu bir hata değildir."),
    ("h2", "Bağımsızlık"),
    ("p", "Her yorumu diğerlerinden bağımsız değerlendirin; benzer metinleri tutarlı hâle getirmeye çalışmayın. Diğer etiketleyiciyle etiketleme bitene kadar kararlarınızı paylaşmayın veya tartışmayın. Bir yoruma 30 saniyeden fazla takılırsanız en uygun etiketi seçip 'Not' sütununa kısa bir açıklama yazabilirsiniz."),
    ("h2", "Doldurma örneği"),
    ("t", "Metin|'Kokusu güzel ama hiç kalıcı değil, para çöpe.'"),
    ("t", "Etiket|Negatif   (Not: boş bırakılabilir)"),
]


def build(path, order_seed):
    df = items.sample(frac=1, random_state=order_seed).reset_index(drop=True)
    wb = Workbook()
    g = wb.active; g.title = "Kılavuz"
    g.column_dimensions["A"].width = 34; g.column_dimensions["B"].width = 110
    r = 1
    for kind, txt in GUIDE:
        txt = txt.replace("{n}", str(len(df)))
        if kind == "h1":
            g.cell(r, 1, txt).font = Font(name=F, size=16, bold=True); r += 2; continue
        if kind == "h2":
            r += 1; g.cell(r, 1, txt).font = Font(name=F, size=12, bold=True); r += 1; continue
        if kind == "progress":
            g.cell(r, 1, "Tamamlanan madde").font = Font(name=F, bold=True)
            c = g.cell(r, 2, f"=COUNTA(Etiketleme!C2:C{len(df)+1})&\" / {len(df)}\"")
            c.font = Font(name=F, bold=True); r += 1; continue
        if kind == "t":
            a, b = txt.split("|", 1)
            g.cell(r, 1, a).font = Font(name=F, bold=True)
            c = g.cell(r, 2, b)
        else:
            g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            c = g.cell(r, 1, txt)
        c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical="top")
        g.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        g.row_dimensions[r].height = max(15, 15 * (len(txt) // 120 + 1))
        r += 1

    s = wb.create_sheet("Etiketleme")
    for j, (h, w) in enumerate([("No", 9), ("Metin", 100), ("Etiket", 22), ("Not (isteğe bağlı)", 40)], 1):
        c = s.cell(1, j, h); c.font = Font(name=F, bold=True); c.fill = HEAD
        s.column_dimensions[c.column_letter].width = w
    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Geçersiz etiket",
                        error="Lütfen listeden bir etiket seçin.")
    s.add_data_validation(dv)
    for i, row in df.iterrows():
        rr = i + 2
        s.cell(rr, 1, row.item_id).font = Font(name=F)
        c = s.cell(rr, 2, row.text); c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical="top")
        e = s.cell(rr, 3); e.fill = YELLOW; e.font = Font(name=F)
        s.cell(rr, 4).font = Font(name=F)
        for j in range(1, 5):
            s.cell(rr, j).border = Border(bottom=thin)
        s.row_dimensions[rr].height = min(15 * (len(str(row.text)) // 95 + 1), 150)
    dv.add(f"C2:C{len(df)+1}")
    s.freeze_panes = "A2"
    wb.active = 1
    wb.save(path)
    return df


a = build("share/Etiketleme_A.xlsx", 11)
b = build("share/Etiketleme_B.xlsx", 22)
assert set(a.item_id) == set(b.item_id) and (a.item_id.values != b.item_id.values).any()
print(len(a), "madde; sıralar farklı")
